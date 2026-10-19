from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def usage_text(record: dict[str, Any]) -> str:
    usage = record.get("metrics", {}).get("usage_normalized", {})
    fields = [
        ("输入 Token", usage.get("input_tokens")),
        ("输出 Token", usage.get("output_tokens")),
        ("合计 Token", usage.get("total_tokens")),
        ("缓存 Token", usage.get("cached_input_tokens")),
        ("思考 Token", usage.get("reasoning_tokens")),
    ]
    return "\n".join(f"{label}：{value if value is not None else '未提供'}" for label, value in fields)


def error_text(record: dict[str, Any]) -> str:
    response = record.get("response", {})
    parts = [
        response.get("error"),
        response.get("parse_error"),
        "；".join(response.get("validation_errors") or []),
    ]
    return "；".join(str(part) for part in parts if part) or "模型调用或结果校验未通过"


def overall_content(record: dict[str, Any]) -> str:
    if not record.get("success"):
        return f"总体测试失败：{error_text(record)}"
    data = record.get("response", {}).get("parsed_json") or {}
    scene = data.get("scene") or {}
    lines = []
    if scene.get("context_zh"):
        lines.append(f"场景描述：{scene['context_zh']}")
    if scene.get("caption_en"):
        lines.append(f"英文场景表达：{scene['caption_en']}")
    if scene.get("caption_zh"):
        lines.append(f"中文场景表达：{scene['caption_zh']}")
    words = []
    for index, item in enumerate(data.get("items") or [], start=1):
        en = item.get("label_en") or item.get("name_en") or ""
        zh = item.get("label_zh") or item.get("name_zh") or ""
        words.append(f"{index}. {en}｜{zh}".rstrip("｜"))
    if words:
        lines.append("学习词汇：\n" + "\n".join(words))
    return "\n\n".join(lines) or "总体结果中没有可汇总内容"


def item_content(record: dict[str, Any], item: dict[str, Any] | None) -> str:
    item_id = record.get("test_context", {}).get("item_id") or ""
    en = ""
    zh = ""
    if item:
        item_id = item.get("id") or item_id
        en = item.get("label_en") or item.get("name_en") or ""
        zh = item.get("label_zh") or item.get("name_zh") or ""
    title = "｜".join(part for part in [str(item_id), en, zh] if part)
    if not record.get("success"):
        prefix = f"学习内容：{title}\n\n" if title else ""
        return prefix + f"详情测试失败：{error_text(record)}"
    data = record.get("response", {}).get("parsed_json") or {}
    lines = []
    if title:
        lines.append(f"学习内容：{title}")
    if data.get("meaning_zh"):
        lines.append(f"释义：{data['meaning_zh']}")
    for index, example in enumerate(data.get("examples") or [], start=1):
        if not isinstance(example, dict):
            continue
        en_example = example.get("en", "")
        zh_example = example.get("zh", "")
        lines.append(f"例句 {index}：{en_example}\n中文译文：{zh_example}")
    if data.get("usage_note_zh"):
        lines.append(f"用法提示：{data['usage_note_zh']}")
    return "\n\n".join(lines) or "详情结果中没有可汇总内容"


def result_rows(run_dir: Path) -> list[list[Any]]:
    overall_files = sorted(run_dir.glob("*-总体结果.json"))
    if len(overall_files) != 1:
        raise ValueError("目录中必须有且仅有一个“模型名-总体结果.json”")
    overall = load_json(overall_files[0])
    parsed = overall.get("response", {}).get("parsed_json") or {}
    items = parsed.get("items") if isinstance(parsed, dict) else []
    item_map = {
        str(item.get("id")): item
        for item in (items or [])
        if isinstance(item, dict) and item.get("id") is not None
    }
    rows: list[list[Any]] = [[
        overall_content(overall),
        overall.get("metrics", {}).get("wall_clock_seconds"),
        usage_text(overall),
    ]]
    item_files = list(run_dir.glob("*-item*.json"))
    item_files.sort(key=lambda p: int(re.search(r"-item(\d+)\.json$", p.name).group(1)) if re.search(r"-item(\d+)\.json$", p.name) else 10**9)
    for path in item_files:
        record = load_json(path)
        item_id = str(record.get("test_context", {}).get("item_id") or "")
        rows.append([
            item_content(record, item_map.get(item_id)),
            record.get("metrics", {}).get("wall_clock_seconds"),
            usage_text(record),
        ])
    return rows


def generate_report(run_dir: Path) -> Path:
    from copy import copy
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Side

    rows = result_rows(run_dir)
    output = run_dir / "结果汇总.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "人工审核"
    ws.sheet_view.showGridLines = False
    ws.append(["详情内容", "时延（秒）", "Token 消耗"])
    for row in rows:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="0B766E")
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    body_font = Font(name="微软雅黑", color="1F2933", size=10)
    line = Side(style="thin", color="D7E0E5")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        border = copy(cell.border)
        border.bottom = line
        cell.border = border
    ws.row_dimensions[1].height = 28

    for row in range(2, ws.max_row + 1):
        for col in range(1, 4):
            cell = ws.cell(row, col)
            cell.font = body_font
            cell.alignment = Alignment(
                horizontal="left" if col in (1, 3) else "center",
                vertical="top" if col in (1, 3) else "center",
                wrap_text=True,
            )
        text_lines = str(ws.cell(row, 1).value or "").count("\n") + 1
        ws.row_dimensions[row].height = min(250, max(72, 17 * text_lines + 12))

    ws.column_dimensions["A"].width = 82
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 18
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{ws.max_row}"
    wb.save(output)

    check = load_workbook(output, data_only=False)
    check_ws = check["人工审核"]
    if check_ws.max_column != 3 or check_ws.max_row != len(rows) + 1:
        raise ValueError("Excel 回读结构与预期不一致")
    if [check_ws.cell(1, c).value for c in range(1, 4)] != ["详情内容", "时延（秒）", "Token 消耗"]:
        raise ValueError("Excel 表头回读不一致")
    return output


def discover_run_dirs(root: Path) -> list[Path]:
    candidates = []
    if len(list(root.glob("*-总体结果.json"))) == 1:
        candidates.append(root)
    for path in root.rglob("*"):
        if path.is_dir() and len(list(path.glob("*-总体结果.json"))) == 1:
            candidates.append(path)
    return sorted(set(candidates))


def generate_reports(root: Path) -> dict[str, Any]:
    run_dirs = discover_run_dirs(root)
    generated = []
    failed = []
    for run_dir in run_dirs:
        try:
            report = generate_report(run_dir)
            generated.append(str(report))
        except Exception as exc:
            failed.append({"directory": str(run_dir), "error": f"{type(exc).__name__}: {exc}"})
    return {"generated": generated, "failed": failed, "run_directory_count": len(run_dirs)}
