#!/usr/bin/env python3
"""Integration test for batch runs, retry context, and per-image Excel reports."""

from __future__ import annotations

import json
import os
import signal
import subprocess
import sys
import tempfile
import time
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PIXEL = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZmxsAAAAASUVORK5CYII="


def wait_line(proc: subprocess.Popen[str], timeout: float = 5.0) -> str:
    started = time.time()
    while time.time() - started < timeout:
        line = proc.stdout.readline().strip()  # type: ignore[union-attr]
        if line:
            return line
    raise TimeoutError("process did not print startup line")


def post(url: str, payload: dict) -> dict:
    req = urllib.request.Request(
        url,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=15) as response:
        return json.loads(response.read().decode("utf-8"))


def stop_process(proc: subprocess.Popen[str] | None) -> None:
    if proc is None or proc.poll() is not None:
        return
    os.killpg(proc.pid, signal.SIGTERM)
    try:
        proc.wait(timeout=3)
    except subprocess.TimeoutExpired:
        os.killpg(proc.pid, signal.SIGKILL)
        proc.wait(timeout=3)


def main() -> None:
    upstream = subprocess.Popen(
        [sys.executable, str(ROOT / "tests" / "mock_upstream.py")],
        stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, start_new_session=True,
    )
    app = None
    try:
        upstream_port = int(wait_line(upstream))
        with tempfile.TemporaryDirectory() as tmp:
            app = subprocess.Popen(
                [sys.executable, str(ROOT / "server.py"), "--port", "0", "--results-dir", tmp],
                stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, start_new_session=True,
            )
            app_port = int(wait_line(app).rsplit(":", 1)[1])
            api = f"http://127.0.0.1:{app_port}"
            config = {
                "provider": "mock",
                "base_url": f"http://127.0.0.1:{upstream_port}/v1",
                "model": "mock-model",
                "model_label": "seed-2.1-lite",
                "api_key": "secret-not-persisted",
                "timeout_seconds": 10,
                "extra_body": {"response_format": {"type": "json_object"}},
                "user_profile": "测试用户资料",
            }
            batch_id = post(f"{api}/api/batch/start", {})["batch_id"]

            success_prepared = post(f"{api}/api/batch/prepare-image", {
                "batch_id": batch_id,
                "model_label": config["model_label"],
                "image_name": "success.png",
                "image_data_url": PIXEL,
            })
            success_id = success_prepared["test_run_id"]
            overall = post(f"{api}/api/run", {
                **config,
                "phase": "overall",
                "test_run_id": success_id,
                "allow_existing_overall": True,
                "batch_id": batch_id,
                "prompt": "总体提示词",
                "image_data_url": PIXEL,
                "image_original_name": "success.png",
            })
            assert overall["result"]["success"] is True
            parsed = overall["result"]["response"]["parsed_json"]
            assert len(parsed["items"]) == 1
            item = parsed["items"][0]
            detail = post(f"{api}/api/run", {
                **config,
                "phase": "item",
                "test_run_id": success_id,
                "batch_id": batch_id,
                "item_index": 1,
                "item_id": item["id"],
                "prompt": "你是英语学习内容编辑，为指定条目生成简短详情。",
                "image_data_url": PIXEL,
                "image_original_name": "success.png",
            })
            assert detail["result"]["success"] is True

            failed_prepared = post(f"{api}/api/batch/prepare-image", {
                "batch_id": batch_id,
                "model_label": config["model_label"],
                "image_name": "failure.png",
                "image_data_url": PIXEL,
            })
            failed_id = failed_prepared["test_run_id"]
            failed = post(f"{api}/api/run", {
                **config,
                "phase": "overall",
                "test_run_id": failed_id,
                "allow_existing_overall": True,
                "batch_id": batch_id,
                "prompt": "FORCE_OVERALL_FAILURE",
                "image_data_url": PIXEL,
                "image_original_name": "failure.png",
            })
            assert failed["result"]["success"] is False

            summary = post(f"{api}/api/batch/finalize", {"batch_id": batch_id})
            assert summary["test_count"] == 2
            assert summary["success_count"] == 1
            assert summary["failure_count"] == 1
            assert summary["failures"][0]["test_run_id"] == failed_id

            success_dir = Path(tmp) / success_id
            failed_dir = Path(tmp) / failed_id
            assert len(list(success_dir.glob("原图-*"))) == 1
            assert len(list(failed_dir.glob("原图-*"))) == 1
            assert len(list(success_dir.glob("*.json"))) == 2
            assert len(list(failed_dir.glob("*.json"))) == 1

            context = post(f"{api}/api/run-context", {"test_run_id": failed_id})
            assert context["test_run_id"] == failed_id
            assert context["image_data_url"].startswith("data:image/png;base64,")
            assert context["overall_record"]["success"] is False
            serialized = json.dumps(context, ensure_ascii=False)
            assert "secret-not-persisted" not in serialized

            manual_save = post(f"{api}/api/run-context/save-overall", {
                "test_run_id": failed_id,
                "overall_json": parsed,
            })
            assert manual_save["updated"] is True
            manual_record = json.loads((failed_dir / "seed-2.1-lite-总体结果.json").read_text(encoding="utf-8"))
            assert manual_record["success"] is True
            assert manual_record["test_context"]["manual_override"] is True
            assert len(list((failed_dir / "archive").glob("*.json"))) == 1

            retry = post(f"{api}/api/run", {
                **config,
                "phase": "overall",
                "test_run_id": failed_id,
                "allow_existing_overall": True,
                "batch_id": batch_id,
                "prompt": "重试总体提示词",
                "image_data_url": PIXEL,
                "image_original_name": "failure.png",
            })
            assert retry["result"]["success"] is True
            assert retry["test_run_id"] == failed_id
            assert len(list((failed_dir / "archive").glob("*.json"))) == 2

            reports = post(f"{api}/api/reports/generate", {"path": batch_id})
            assert reports["run_directory_count"] == 2
            assert len(reports["generated_relative"]) == 2
            assert all((Path(tmp) / path).is_file() for path in reports["generated_relative"])

            from openpyxl import load_workbook
            workbook = load_workbook(success_dir / "结果汇总.xlsx", data_only=True)
            sheet = workbook["人工审核"]
            assert sheet.max_row == 3
            assert sheet.max_column == 3
            assert "场景描述：" in sheet["A2"].value
            assert "学习词汇：" in sheet["A2"].value
            assert "释义：" in sheet["A3"].value
            assert "例句 1：" in sheet["A3"].value
            print("PASS: batch folders, images, failure recovery, retry overwrite, Excel summaries")
    finally:
        stop_process(app)
        stop_process(upstream)


if __name__ == "__main__":
    main()
